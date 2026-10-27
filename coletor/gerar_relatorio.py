"""Gerador do relatorio diario de prazos — arquivo unico.

Montado a partir do pacote `coletor` para rodar sem instalacao: a tarefa
agendada baixa este arquivo e executa. Nao altere aqui: altere no pacote
e gere de novo com `python scripts/monta_gerador.py`.

    python gerar_relatorio.py --json publicacoes.json --saida pasta \
        --advogado "Nome" --oab "229618/RJ"

Regras: docs/regras-classificacao.md, decididas em 27/10/2026.
"""
import argparse
import datetime as dt
import html
import json
import re
import unicodedata
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (KeepTogether, Paragraph, SimpleDocTemplate,
                                Spacer, Table, TableStyle)

# --- Prioridades, na ordem em que aparecem no relatorio ---
SIGILO = "SIGILO/ATENÇÃO"
URGENTE = "URGENTE"
ALTA = "ALTA"
MEDIA = "MÉDIA"
BAIXA = "BAIXA"
ORDEM_PRIORIDADE = [SIGILO, URGENTE, ALTA, MEDIA, BAIXA]

CORES = {
    SIGILO: "7030A0",
    URGENTE: "C00000",
    ALTA: "E36C09",
    MEDIA: "8EA9DB",
    BAIXA: "A9D18E",
}

# Quantos dias uteis antes do prazo final fica o prazo interno do escritorio
FOLGA_ESCRITORIO = 3

# Como contar os N dias uteis a partir do inicio.
#   True  -> o dia de inicio conta como dia 1 (final = inicio + N-1 dias uteis)
#   False -> a contagem comeca no dia util seguinte ao inicio (final = inicio + N)
# O texto da automacao ("soma de N dias uteis a partir do inicio") admite as
# duas leituras, e o historico do escritorio usa ora uma, ora outra. Adotada a
# primeira, que segue o art. 224 do CPC: exclui-se o dia do comeco (a
# publicacao) e inclui-se o do vencimento. Trocar aqui muda todo o sistema.
INICIO_CONTA_COMO_DIA_1 = True

RX_DIAS = re.compile(r"\b(?:em|prazo de)\s+(\d{1,3})\s+dias\b", re.IGNORECASE)
RX_DATA_BR = re.compile(r"\b(\d{2})/(\d{2})/(\d{4})\b")
RX_AUDIENCIA = re.compile(r"audi[êe]ncia", re.IGNORECASE)

TERMOS_ANDAMENTO = ("arquivamento", "sem saldo", "nomeacao de perito",
                    "lista de distribuicao", "distribuido")


RX_TAG = re.compile(r"<[^>]+>")
RX_BR = re.compile(r"<\s*br\s*/?\s*>", re.IGNORECASE)
RX_ESPACO = re.compile(r"\s+")


def limpa_texto(bruto):
    """Tira o HTML que o DJEN manda no campo `texto`.

    O teor vem com <br> e entidades HTML. Isso quebra o gerador de PDF e
    atrapalha a busca do prazo, entao a limpeza vale para os dois usos.
    """
    if not bruto:
        return ""
    texto = RX_BR.sub(" ", str(bruto))
    texto = RX_TAG.sub(" ", texto)
    texto = html.unescape(texto)
    # Depois do unescape pode sobrar < ou & solto, que o reportlab recusa
    texto = texto.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return RX_ESPACO.sub(" ", texto).strip()


def _sem_acento(texto):
    return "".join(c for c in unicodedata.normalize("NFD", texto or "")
                   if unicodedata.category(c) != "Mn").lower()


# ---------------------------------------------------------------- datas

def e_dia_util(data):
    """Dia util = segunda a sexta. Feriado nao entra, por decisao."""
    return data.weekday() < 5


def proximo_dia_util(data):
    """Primeiro dia util DEPOIS de `data` (nunca a propria data)."""
    data += dt.timedelta(days=1)
    while not e_dia_util(data):
        data += dt.timedelta(days=1)
    return data


def soma_dias_uteis(data, quantidade):
    """Avanca `quantidade` dias uteis a partir de `data`, sem conta-la."""
    for _ in range(quantidade):
        data = proximo_dia_util(data)
    return data


def subtrai_dias_uteis(data, quantidade):
    """Recua `quantidade` dias uteis a partir de `data`, sem conta-la."""
    for _ in range(quantidade):
        data -= dt.timedelta(days=1)
        while not e_dia_util(data):
            data -= dt.timedelta(days=1)
    return data


def para_data(valor):
    """Le a data da API em qualquer dos formatos observados."""
    if isinstance(valor, dt.datetime):
        return valor.date()
    if isinstance(valor, dt.date):
        return valor
    if not valor:
        return None
    texto = str(valor).strip()[:19]
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%dT%H:%M:%S",
                    "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    return None


def calcula_prazo(data_disponibilizacao, dias):
    """Aplica a contagem vigente. Devolve as quatro datas do fluxo."""
    publicacao = proximo_dia_util(data_disponibilizacao)
    inicio = proximo_dia_util(publicacao)
    passos = dias - 1 if INICIO_CONTA_COMO_DIA_1 else dias
    final = soma_dias_uteis(inicio, passos)
    return {
        "publicacao": publicacao,
        "inicio": inicio,
        "prazo_final": final,
        "prazo_escritorio": subtrai_dias_uteis(final, FOLGA_ESCRITORIO),
    }


# ---------------------------------------------------------- classificacao

def e_trabalhista(sigla_tribunal):
    sigla = (sigla_tribunal or "").strip().upper()
    return sigla.startswith("TRT") or sigla == "TST"


def prioridade_por_dias(dias):
    if dias <= 5:
        return URGENTE
    if dias <= 15:
        return ALTA
    return MEDIA


def classifica(item):
    """Aplica as seis regras, na ordem. A primeira que casar decide.

    Recebe um item cru do DJEN e devolve a classificacao, sem alterar o item.
    """
    texto = limpa_texto(item.get("texto"))
    tribunal = item.get("siglaTribunal") or ""
    disponibilizacao = para_data(item.get("data_disponibilizacao"))

    base = {
        "prioridade": MEDIA,
        "dias": None,
        "publicacao": None,
        "inicio": None,
        "prazo_final": None,
        "prazo_escritorio": None,
        "regra_aplicada": "",
        "observacao": "",
    }

    # Regra 1 - sem teor
    if not texto:
        base.update(prioridade=SIGILO, regra_aplicada="Sem teor na API pública",
                    observacao="sem teor disponível na API pública, possível sigilo")
        return base

    achou = RX_DIAS.search(texto)

    # Regra 3 - trabalhista com prazo: nao calcula, por decisao
    if achou and e_trabalhista(tribunal):
        base.update(prioridade=MEDIA, dias=int(achou.group(1)),
                    regra_aplicada="Justiça do Trabalho — prazo não calculado (decisão de 27/10/2026)",
                    observacao="regra CLT — requer conferência manual")
        return base

    # Regra 2 - prazo em dias, fora da Justica do Trabalho
    if achou:
        dias = int(achou.group(1))
        base.update(prioridade=prioridade_por_dias(dias), dias=dias,
                    regra_aplicada=f"Prazo de {dias} dias úteis identificado no teor")
        if disponibilizacao:
            base.update(calcula_prazo(disponibilizacao, dias))
        else:
            base["observacao"] = "data de disponibilização ausente — prazo não calculado"
        return base

    # Regra 4 - audiencia com data fixa
    if RX_AUDIENCIA.search(texto):
        data = RX_DATA_BR.search(texto)
        if data:
            dia, mes, ano = (int(x) for x in data.groups())
            try:
                base.update(prioridade=MEDIA, prazo_final=dt.date(ano, mes, dia),
                            regra_aplicada="Data fixa de audiência",
                            observacao="data fixa de audiência")
                return base
            except ValueError:
                pass

    # Regra 5 - mero andamento
    plano = _sem_acento(texto)
    if any(termo in plano for termo in TERMOS_ANDAMENTO):
        base.update(prioridade=BAIXA, regra_aplicada="Mero andamento",
                    observacao="sem prazo para o escritório")
        return base

    # Regra 6 - resto
    base.update(prioridade=MEDIA, regra_aplicada="Nenhum prazo identificado",
                observacao="nenhum prazo identificado automaticamente — "
                           "requer leitura manual do teor")
    return base


def classifica_lote(itens):
    """Classifica e deduplica por `id`, preservando a ordem de chegada."""
    vistos = set()
    saida = []
    for item in itens:
        # `or` nao serve aqui: id igual a 0 e valido e seria tratado como ausente
        chave = item.get("id")
        if chave is None:
            chave = item.get("hash")
        if chave is not None and chave in vistos:
            continue
        if chave is not None:
            vistos.add(chave)
        registro = dict(item)
        registro["classificacao"] = classifica(item)
        saida.append(registro)
    return saida


def ordena_por_prioridade(registros):
    posicao = {p: i for i, p in enumerate(ORDEM_PRIORIDADE)}
    return sorted(registros,
                  key=lambda r: (posicao.get(r["classificacao"]["prioridade"], 99),
                                 r["classificacao"]["prazo_final"] or dt.date.max))


AVISO = (
    "Classificação e prazos gerados automaticamente por heurística de texto. "
    "<b>A validação humana é obrigatória.</b> O cálculo não considera feriados "
    "forenses, feriados locais nem o recesso de 20/12 a 20/01. Publicações do "
    "TRT e do TST <b>não têm prazo calculado</b> e dependem de conferência "
    "manual. Todo prazo aqui nasce não confirmado."
)


def _texto(valor):
    return "" if valor is None else str(valor)


def _partes(item):
    destinatarios = item.get("destinatarios") or []
    nomes = []
    for d in destinatarios:
        if isinstance(d, dict):
            nomes.append(d.get("nome") or d.get("nomeParte") or "")
        else:
            nomes.append(str(d))
    nomes = [n for n in nomes if n]
    if not nomes:
        return "—"
    if len(nomes) > 3:
        return "; ".join(nomes[:3]) + f" e mais {len(nomes) - 3}"
    return "; ".join(nomes)


def junta(*partes, sep=" · "):
    """Junta os pedacos nao vazios; evita separador solto."""
    return sep.join(str(x).strip() for x in partes if x and str(x).strip()) or "—"


def _data_br(valor):
    if isinstance(valor, dt.date):
        return valor.strftime("%d/%m/%Y")
    return "—"


def sugestao(classificacao):
    """Sugestao de medida, derivada da regra. Sempre editavel pelo advogado."""
    p = classificacao["prioridade"]
    if p == SIGILO:
        return "Consultar o teor no sistema do tribunal e classificar manualmente"
    if classificacao["dias"] and classificacao["prazo_final"] is None:
        return ("Conferir a contagem manualmente e lançar o prazo "
                f"({classificacao['dias']} dias)")
    if classificacao["prazo_final"] and classificacao["observacao"] == "data fixa de audiência":
        return "Confirmar a pauta e agendar preparação da audiência"
    if classificacao["prazo_final"]:
        return "Protocolar a manifestação até o prazo do escritório"
    if p == BAIXA:
        return "Apenas ciência"
    return "Ler o teor e decidir a medida"


def gera_pdf(registros, coleta_resumo, advogado, oab, caminho, data=None):
    data = data or dt.date.today()
    estilos = getSampleStyleSheet()
    normal = ParagraphStyle("corpo", parent=estilos["Normal"], fontSize=8.5,
                            leading=11, alignment=TA_JUSTIFY)
    rotulo = ParagraphStyle("rotulo", parent=normal, fontSize=7.5,
                            textColor=colors.HexColor("#555555"))
    vermelho = ParagraphStyle("vermelho", parent=normal, fontSize=9,
                              textColor=colors.HexColor("#C00000"))
    titulo = ParagraphStyle("titulo", parent=estilos["Title"], fontSize=16)
    aviso = ParagraphStyle("aviso", parent=normal, fontSize=8,
                           textColor=colors.HexColor("#C00000"),
                           fontName="Helvetica-Oblique")

    doc = SimpleDocTemplate(caminho, pagesize=A4,
                            leftMargin=15*mm, rightMargin=15*mm,
                            topMargin=15*mm, bottomMargin=15*mm,
                            title="Relatório Diário de Publicações e Prazos")
    fluxo = [Paragraph("Relatório Diário de Publicações e Prazos", titulo),
             Paragraph(f"{advogado} — OAB {oab} — {data.strftime('%d/%m/%Y')}",
                       estilos["Normal"]),
             Spacer(1, 4*mm)]

    # --- Estado da coleta, antes de qualquer numero ---
    if not coleta_resumo["completa"]:
        fluxo += [Paragraph(
            f"<b>ATENÇÃO: a coleta falhou.</b> {coleta_resumo['erro']}<br/>"
            "Os números abaixo <b>não</b> representam o período. "
            "Não é possível afirmar que não houve publicações.", aviso),
            Spacer(1, 4*mm)]
    elif not registros:
        fluxo += [Paragraph(
            "Nenhuma publicação encontrada no período consultado "
            f"({coleta_resumo['periodo']['inicio']} a "
            f"{coleta_resumo['periodo']['fim']}). A consulta foi concluída "
            "com sucesso.", normal), Spacer(1, 4*mm)]

    # --- Resumo por prioridade ---
    contagem = {p: 0 for p in ORDEM_PRIORIDADE}
    for r in registros:
        contagem[r["classificacao"]["prioridade"]] += 1
    linhas = [["Prioridade", "Publicações"]]
    estilo_resumo = [
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEEEEE")),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("ALIGN", (1, 0), (1, -1), "CENTER"),
    ]
    for i, p in enumerate(ORDEM_PRIORIDADE, start=1):
        linhas.append([p, str(contagem[p])])
        estilo_resumo += [
            ("BACKGROUND", (0, i), (0, i), colors.HexColor("#" + CORES[p])),
            ("TEXTCOLOR", (0, i), (0, i), colors.white),
        ]
    linhas.append(["Total", str(len(registros))])
    tabela = Table(linhas, colWidths=[60*mm, 25*mm])
    tabela.setStyle(TableStyle(estilo_resumo))
    fluxo += [tabela, Spacer(1, 4*mm), Paragraph(AVISO, aviso), Spacer(1, 6*mm)]

    # --- Cards, agrupados por prioridade ---
    for prioridade in ORDEM_PRIORIDADE:
        do_grupo = [r for r in registros
                    if r["classificacao"]["prioridade"] == prioridade]
        if not do_grupo:
            continue
        cor = colors.HexColor("#" + CORES[prioridade])
        cabecalho = Table([[Paragraph(
            f'<font color="white"><b>{prioridade} — {len(do_grupo)}</b></font>',
            normal)]], colWidths=[180*mm])
        cabecalho.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), cor),
                                       ("LEFTPADDING", (0, 0), (-1, -1), 6)]))
        fluxo += [cabecalho, Spacer(1, 2*mm)]
        for r in do_grupo:
            fluxo.append(_card(r, cor, normal, rotulo, vermelho))
            fluxo.append(Spacer(1, 3*mm))

    doc.build(fluxo)
    return caminho


def _card(registro, cor, normal, rotulo, vermelho):
    c = registro["classificacao"]
    teor = limpa_texto(registro.get("texto"))
    if len(teor) > 700:
        teor = teor[:700] + " […]"

    def par(rot, valor, estilo=None):
        return [Paragraph(rot, rotulo), Paragraph(_texto(valor), estilo or normal)]

    datas = (f"Disponibilização: {_data_br(para_data(registro.get('data_disponibilizacao')))} · "
             f"Publicação: {_data_br(c['publicacao'])} · "
             f"Início: {_data_br(c['inicio'])}")
    if c["prazo_final"]:
        prazo = (f"<b>Prazo final: {_data_br(c['prazo_final'])}</b>")
        if c["prazo_escritorio"]:
            prazo += f" · Prazo do escritório: <b>{_data_br(c['prazo_escritorio'])}</b>"
    else:
        prazo = "<b>Prazo não calculado</b> — " + (c["observacao"] or "ver regra aplicada")

    linhas = [
        par("Processo", registro.get("numeroprocessocommascara")
            or registro.get("numero_processo") or "—"),
        par("Cliente / partes", _partes(registro)),
        par("Tribunal / órgão", junta(registro.get("siglaTribunal"),
                                      registro.get("nomeOrgao"), sep=" — ")),
        par("Tipo de ato", junta(registro.get("tipoComunicacao"),
                                 registro.get("tipoDocumento"))),
        par("Teor", teor or "—"),
        par("Datas", datas),
        par("Prazo", prazo, vermelho),
        par("Regra aplicada", c["regra_aplicada"] or "—"),
        par("Sugestão de medida", sugestao(c)),
    ]
    tabela = Table(linhas, colWidths=[32*mm, 148*mm])
    tabela.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEBEFORE", (0, 0), (0, -1), 2, cor),
        ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#CCCCCC")),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
    ]))
    return KeepTogether(tabela)


COLUNAS = [
    ("Prioridade", 16), ("Processo", 24), ("Cliente/Parte", 34),
    ("Tribunal/Órgão", 30), ("Tipo de ato", 22), ("Teor", 60),
    ("Disponibilização", 15), ("Publicação", 13), ("Início do prazo", 15),
    ("Dias úteis do prazo", 12), ("Prazo final", 13),
    ("Prazo do escritório", 15), ("Regra aplicada", 38),
    ("Sugestão de medida", 34),
]

LEIAME = [
    ("Como ler esta planilha", True),
    ("", False),
    ("Gerada automaticamente a partir do DJEN (API pública do CNJ).", False),
    ("", False),
    ("O que ela NÃO faz:", True),
    ("1. Não considera feriados forenses, feriados locais nem o recesso de 20/12 a 20/01.", False),
    ("2. Não calcula prazo de TRT nem TST. Esses itens vêm sem prazo e precisam de conferência manual.", False),
    ("3. Identifica o número de dias por expressão de texto: não pega \"15 (quinze) dias\" nem prazo citado por artigo de lei.", False),
    ("4. Não guarda histórico: cada dia gera um arquivo novo.", False),
    ("", False),
    ("Toda classificação e todo prazo aqui são sugestões automáticas.", True),
    ("A validação é do advogado. Nenhum prazo nasce confirmado.", True),
    ("", False),
    ("As colunas Publicação, Início, Prazo final e Prazo do escritório usam fórmula WORKDAY:", False),
    ("mudando a coluna \"Dias úteis do prazo\", as datas se recalculam sozinhas.", False),
    ("WORKDAY considera só segunda a sexta — nenhum feriado é informado a ela.", False),
]


def gera_xlsx(registros, coleta_resumo, advogado, oab, caminho, data=None):
    data = data or dt.date.today()
    wb = Workbook()

    # --- Aba Leia-me ---
    ws = wb.active
    ws.title = "Leia-me"
    ws.column_dimensions["A"].width = 115
    ws["A1"] = f"Controle de prazos — {advogado} — OAB {oab} — {data.strftime('%d/%m/%Y')}"
    ws["A1"].font = Font(bold=True, size=13)
    linha = 3
    if not coleta_resumo["completa"]:
        ws[f"A{linha}"] = f"ATENÇÃO: a coleta falhou. {coleta_resumo['erro']}"
        ws[f"A{linha}"].font = Font(bold=True, color="C00000")
        ws[f"A{linha + 1}"] = ("Os dados abaixo não representam o período. "
                               "Não é possível afirmar que não houve publicações.")
        ws[f"A{linha + 1}"].font = Font(color="C00000")
        linha += 3
    for texto, negrito in LEIAME:
        ws[f"A{linha}"] = texto
        ws[f"A{linha}"].font = Font(bold=negrito)
        ws[f"A{linha}"].alignment = Alignment(wrap_text=True)
        linha += 1

    # --- Aba de dados ---
    ws = wb.create_sheet("Prazos")
    for i, (nome, largura) in enumerate(COLUNAS, start=1):
        celula = ws.cell(row=1, column=i, value=nome)
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="404040")
        celula.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.freeze_panes = "A2"

    for i, registro in enumerate(ordena_por_prioridade(registros), start=2):
        c = registro["classificacao"]
        disponibilizacao = para_data(registro.get("data_disponibilizacao"))
        teor = (limpa_texto(registro.get("texto"))
                .replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">"))

        ws.cell(row=i, column=1, value=c["prioridade"]).fill = PatternFill(
            "solid", fgColor=CORES[c["prioridade"]])
        ws.cell(row=i, column=1).font = Font(bold=True, color="FFFFFF")
        ws.cell(row=i, column=2, value=registro.get("numeroprocessocommascara")
                or registro.get("numero_processo") or "")
        ws.cell(row=i, column=3, value=_partes(registro))
        ws.cell(row=i, column=4, value=junta(registro.get("siglaTribunal"),
                                             registro.get("nomeOrgao"), sep=" — "))
        ws.cell(row=i, column=5, value=junta(registro.get("tipoComunicacao"),
                                             registro.get("tipoDocumento")))
        ws.cell(row=i, column=6, value=teor[:2000])
        ws.cell(row=i, column=7, value=disponibilizacao).number_format = "DD/MM/YYYY"
        ws.cell(row=i, column=10, value=c["dias"])

        # As datas calculadas saem como formula, para recalcularem sozinhas.
        # Nos casos sem prazo (sigilo, trabalhista, mero andamento) ficam vazias.
        if disponibilizacao and c["dias"] and c["prazo_final"]:
            passos = c["dias"] - 1 if INICIO_CONTA_COMO_DIA_1 else c["dias"]
            ws.cell(row=i, column=8, value=f"=WORKDAY(G{i},1)")
            ws.cell(row=i, column=9, value=f"=WORKDAY(H{i},1)")
            ws.cell(row=i, column=11, value=f"=WORKDAY(I{i},{passos})")
            ws.cell(row=i, column=12,
                    value=f"=WORKDAY(K{i},-{FOLGA_ESCRITORIO})")
            for coluna in (8, 9, 11, 12):
                ws.cell(row=i, column=coluna).number_format = "DD/MM/YYYY"
            ws.cell(row=i, column=12).font = Font(bold=True, color="C00000")
        elif c["prazo_final"]:  # data fixa de audiencia: valor, nao formula
            ws.cell(row=i, column=11, value=c["prazo_final"]).number_format = "DD/MM/YYYY"

        ws.cell(row=i, column=13, value=c["regra_aplicada"]
                + (f" — {c['observacao']}" if c["observacao"] else ""))
        ws.cell(row=i, column=14, value=sugestao(c))
        for coluna in (3, 4, 5, 6, 13, 14):
            ws.cell(row=i, column=coluna).alignment = Alignment(
                wrap_text=True, vertical="top")

    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{max(1, len(registros) + 1)}"
    wb.save(caminho)
    return caminho


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--json", required=True,
                    help="JSON com itens no formato do DJEN, ou {items:[...]}")
    ap.add_argument("--saida", default=".")
    ap.add_argument("--advogado", default="")
    ap.add_argument("--oab", default="")
    ap.add_argument("--data", default=None, help="AAAA-MM-DD; padrao: hoje")
    ap.add_argument("--coleta-incompleta", default=None,
                    help="motivo da falha; faz o relatorio avisar em vermelho")
    args = ap.parse_args()

    bruto = json.loads(Path(args.json).read_text(encoding="utf-8-sig"))
    itens = bruto.get("items", bruto) if isinstance(bruto, dict) else bruto
    data = dt.date.fromisoformat(args.data) if args.data else dt.date.today()

    resumo = {"completa": args.coleta_incompleta is None,
              "erro": args.coleta_incompleta,
              "periodo": {"inicio": str(data - dt.timedelta(days=1)),
                          "fim": str(data)}}
    registros = ordena_por_prioridade(classifica_lote(itens))
    pasta = Path(args.saida)
    pasta.mkdir(parents=True, exist_ok=True)
    pdf = pasta / f"relatorio_prazos_{data.isoformat()}.pdf"
    xlsx = pasta / f"controle_prazos_{data.isoformat()}.xlsx"
    gera_pdf(registros, resumo, args.advogado, args.oab, str(pdf), data)
    gera_xlsx(registros, resumo, args.advogado, args.oab, str(xlsx), data)

    contagem = {}
    for r in registros:
        p = r["classificacao"]["prioridade"]
        contagem[p] = contagem.get(p, 0) + 1
    print(json.dumps({
        "publicacoes": len(registros),
        "por_prioridade": contagem,
        "com_prazo_no_texto_sem_calculo": sum(
            1 for r in registros if r["classificacao"]["dias"]
            and r["classificacao"]["prazo_final"] is None),
        "atos_distintos": len({r.get("link") for r in registros if r.get("link")}),
        "pdf": str(pdf), "xlsx": str(xlsx)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
